import streamlit as st
import pandas as pd
import sqlite3
import time as time_module # 防止和之前的time冲突, 起个别名
import requests
import io
from openpyxl import Workbook
import plotly.express as px

api_key = st.secrets["api_key"]
url_api = "https://api.deepseek.com/chat/completions"

st.title("🤖 AI智能数据分析助手")

# ===== 新增: 侧边栏历史记录 =====
with st.sidebar:
    st.header("历史分析记录")
    try:
        conn = sqlite3.connect('D:\\AI分析历史记录.db')
        cursor = conn.cursor()
        cursor.execute("SELECT time, filename FROM history ORDER BY id DESC LIMIT 10")
        rows = cursor.fetchall()
        if rows:
            for row in rows:
                st.write(f"**{row[0]}** - {row[1]}")
        else:
            st.write("暂无记录")
        conn.close()
    except:
        st.write("等待第一条记录...")
# ===================================

# 增加密码验证
password = st.text_input("请输入使用密码: ", type="password")
if password != st.secrets["password"]:
    st.warning("请输入正确的密码才能使用! ")
    st.stop()  # 密码不对, 直接停止运行

# 1. 上传文件 (网页上的上传按钮)
uploaded_file = st.file_uploader("请上传你的Excel表格", type=["xlsx"])

# 2. 输入要求
question = st.text_input("请输入你的分析要求 (不填则自动总结) : ")

# 3. 开始分析按钮
if st.button("开始分析"):
    if uploaded_file is None:
        st.warning("请先上传Excel文件! ")
    else:
        try:
            st.write("正在读取文件并发送给AI, 请稍候...")
            df = pd.read_excel(uploaded_file)
            table_text = df.head(10).to_string()

            if not question:
                question = "请总结业务情况, 并指出异常点"

            prompt = f"你是一个AI资深数据分析师。请根据以下表格数据, {question}。\n\n表格数据: \n{table_text}"

            headers_api = {
                "Authorization": f"Bearer {api_key}",
                "Content-Type": "application/json"
            }
            data_api = {
                "model": "deepseek-chat",
                "messages": [{"role": "user", "content": prompt}]
            }

            response_api = requests.post(url_api, headers=headers_api, json=data_api)

            if "choices" in response_api.json():
                result = response_api.json()["choices"][0]["message"]["content"]
                st.success("分析完成! ")
                st.text_area("AI分析报告如下: ", result, height=300)

                st.write("数据可视化: ")
                numeric_cols = df.select_dtypes(include=['int64', 'float64']).columns

                if len(numeric_cols) > 0:
                    fig = px.bar(df, x=df.columns[0], y=numeric_cols[0], title="数据图表分析")
                    st.plotly_chart(fig, use_container_width=True)
                else:
                    st.info("表格中没有可以绘制的数值数据")
                # =======================================

                # 将结果转成 Excel 格式
                wb = Workbook()
                ws = wb.active
                ws['A1'] = "AI数据分析报告"
                ws['A2'] = result
                buffer = io.BytesIO()
                wb.save(buffer)
                buffer.seek(0)
                
                st.download_button("点击下载Excel报告", buffer, file_name="AI分析报告.xlsx")
            else:
                st.error(f"AI报错: {response_api.json()}")
        except Exception as e:
            st.error(f"运行出错: {e}")